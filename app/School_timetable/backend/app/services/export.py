from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _grid(days: list, periods: list, cell_fn):
    header = ["Period"] + [d.name for d in days]
    rows = [header]
    for p in periods:
        if p.is_break:
            rows.append([p.name] + ["Break"] * len(days))
            continue
        start = p.start_time.strftime("%H:%M") if p.start_time else ""
        row = [f"{p.name}\n{start}"]
        for d in days:
            row.append(cell_fn(d.id, p.id) or "")
        rows.append(row)
    return rows


def _style_excel_sheet(ws, rows: list[list[str]]) -> None:
    header_fill = PatternFill("solid", fgColor="1E4D3A")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="C4A35A"),
        right=Side(style="thin", color="C4A35A"),
        top=Side(style="thin", color="C4A35A"),
        bottom=Side(style="thin", color="C4A35A"),
    )
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in rows:
        ws.append([c.replace("\n", "\n") if isinstance(c, str) else c for c in r])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap
        cell.border = thin
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="Calibri", size=10)
    ws.column_dimensions["A"].width = 16
    for idx in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 18
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 36


def excel_sheets(sheets: list[tuple[str, list, list, object]]) -> bytes:
    """sheets: (title, days, periods, cell_fn)"""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for title, days, periods, cell_fn in sheets:
        name = title[:31] or "Sheet"
        base = name
        n = 2
        while name in used:
            name = f"{base[:28]}_{n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(name)
        rows = _grid(days, periods, cell_fn)
        _style_excel_sheet(ws, rows)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_timetable(title: str, days, periods, cell_fn) -> bytes:
    return excel_sheets([(title, days, periods, cell_fn)])


def _pdf_table(data):
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e4d3a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c4a35a")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#f4efe6")),
                ("ROWBACKGROUNDS", (1, 1), (-1, -1), [colors.white, colors.HexColor("#fbf7f0")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def pdf_sheets(doc_title: str, subtitle: str, sheets: list[tuple[str, list, list, object]]) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story = []
    for i, (title, days, periods, cell_fn) in enumerate(sheets):
        if i:
            story.append(PageBreak())
        story.append(Paragraph(title, styles["Title"]))
        story.append(Paragraph(subtitle, styles["Normal"]))
        story.append(Spacer(1, 12))
        story.append(_pdf_table(_grid(days, periods, cell_fn)))
    if not story:
        story.append(Paragraph(doc_title, styles["Title"]))
    doc.build(story)
    return buf.getvalue()


def pdf_timetable(title: str, subtitle: str, days, periods, cell_fn) -> bytes:
    return pdf_sheets(title, subtitle, [(title, days, periods, cell_fn)])
